"""Gera o modelo de ficha de entrega de EPI para baixar: PDF (A4) e planilha.

Saída: public/modelos/ficha-de-entrega-de-epi.pdf e .xlsx, servidos pelo
artigo /conhecimento/ficha-de-entrega-de-epi-o-que-precisa-constar/.

Nada é escrito à mão aqui. O texto do termo sai do próprio artigo (o bloco
que começa em "Declaro que recebi"), e o endereço do site sai de
src/config/empresa.ts. Se o termo mudar no artigo, rodar de novo; o arquivo
para baixar nunca fica dizendo outra coisa que a página.

O que o modelo não tem, de propósito: número de CA de exemplo (CA inventado
é a regra 1 do projeto), nome de empresa preenchido e qualquer frase que
sugira ser modelo oficial. A NR-6 não publica um.

Dependências: pip install openpyxl reportlab
Uso, da raiz do repositório: python3 docs/ferramentas/gerar-modelo-ficha-epi.py
"""
import os
import re
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from reportlab.lib.colors import HexColor
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.pdfgen import canvas

SLUG = 'ficha-de-entrega-de-epi-o-que-precisa-constar'
SAIDA = 'public/modelos'
NOME_ARQUIVO = 'ficha-de-entrega-de-epi'
COLUNAS = ['Data', 'Equipamento (descrição, tamanho)', 'Nº do CA', 'Qtd.', 'Motivo', 'Assinatura']
MOTIVOS = ['Primeira entrega', 'Desgaste', 'Dano', 'Extravio', 'Troca de modelo']
LINHAS_PDF = 16
LINHAS_XLSX = 40


def ler_fontes():
    artigos = open('src/content/artigos.ts', encoding='utf-8').read()
    i = artigos.index(f"slug: '{SLUG}'")
    fim = artigos.find('\n  {\n    slug:', i + 10)
    bloco = artigos[i:fim if fim > 0 else None]
    m = re.search(r"'(Declaro que recebi[^']+)'", bloco)
    if not m:
        sys.exit('Termo não encontrado no artigo. O bloco começa em "Declaro que recebi"?')
    empresa = open('src/config/empresa.ts', encoding='utf-8').read()
    site = re.search(r"\bsite: '(https://[^']+)'", empresa)
    if not site:
        sys.exit('Campo site não encontrado em src/config/empresa.ts.')
    dominio = site.group(1).removeprefix('https://').rstrip('/')
    return m.group(1), dominio


def rodape(dominio):
    return (f'Modelo gratuito de {dominio}. Não é modelo oficial: a NR-6 exige o registro da entrega, '
            'mas não publica formulário. Valide o termo com o responsável pela segurança do trabalho da empresa.')


# ----------------------------------------------------------------- PDF

TINTA = HexColor('#1a1a1a')
CINZA = HexColor('#6b6b6b')
LINHA = HexColor('#9a9a9a')
FUNDO = HexColor('#f1efea')
VERMELHO = HexColor('#c8102e')


def quebrar(c, texto, fonte, tam, largura):
    palavras, linhas, atual = texto.split(), [], ''
    for p in palavras:
        teste = f'{atual} {p}'.strip()
        if c.stringWidth(teste, fonte, tam) <= largura:
            atual = teste
        else:
            linhas.append(atual)
            atual = p
    if atual:
        linhas.append(atual)
    return linhas


def gerar_pdf(termo, dominio, caminho):
    W, H = A4
    m = 14 * mm
    larg = W - 2 * m
    c = canvas.Canvas(caminho, pagesize=A4)
    c.setTitle('Ficha de controle de entrega de EPI')
    c.setAuthor(dominio)
    c.setSubject('Modelo de ficha de entrega de equipamento de proteção individual')

    y = H - m
    c.setFillColor(VERMELHO)
    c.rect(m, y - 2, 18 * mm, 2.2, stroke=0, fill=1)
    c.setFillColor(TINTA)
    c.setFont('Helvetica-Bold', 16)
    y -= 9 * mm
    c.drawString(m, y, 'Ficha de controle de entrega de EPI')
    c.setFont('Helvetica', 9)
    c.setFillColor(CINZA)
    y -= 5 * mm
    c.drawString(m, y, 'Registro do fornecimento de equipamento de proteção individual ao trabalhador.')

    # Cabeçalho: campos com linha para escrever.
    c.setFillColor(TINTA)
    c.setStrokeColor(LINHA)
    c.setLineWidth(0.6)
    linhas_campos = [
        [('Empresa', 1.0)],
        [('Trabalhador', 0.68), ('Matrícula', 0.32)],
        [('Função', 0.5), ('Setor', 0.5)],
        [('Admissão', 0.31), ('Nº do calçado', 0.23), ('Tam. da luva', 0.23), ('Tam. da roupa', 0.23)],
    ]
    y -= 9 * mm
    for campos in linhas_campos:
        x = m
        for rotulo, frac in campos:
            w = larg * frac
            c.setFont('Helvetica', 7.5)
            c.setFillColor(CINZA)
            c.drawString(x, y + 1.6 * mm, rotulo.upper())
            c.line(x, y - 4.2 * mm, x + w - 3 * mm, y - 4.2 * mm)
            x += w
        y -= 11 * mm

    # Termo de recebimento.
    linhas_termo = quebrar(c, termo, 'Helvetica', 8.6, larg - 8 * mm)
    alt = len(linhas_termo) * 3.9 * mm + 21 * mm
    y -= 1 * mm
    c.setFillColor(FUNDO)
    c.rect(m, y - alt, larg, alt, stroke=0, fill=1)
    c.setFillColor(TINTA)
    c.setFont('Helvetica-Bold', 9)
    ty = y - 6 * mm
    c.drawString(m + 4 * mm, ty, 'Termo de recebimento e responsabilidade')
    c.setFont('Helvetica', 8.6)
    ty -= 5.2 * mm
    for ln in linhas_termo:
        c.drawString(m + 4 * mm, ty, ln)
        ty -= 3.9 * mm
    ty -= 4.5 * mm
    c.setStrokeColor(LINHA)
    c.line(m + 4 * mm, ty, m + larg * 0.62, ty)
    c.line(m + larg * 0.68, ty, m + larg - 4 * mm, ty)
    c.setFont('Helvetica', 7)
    c.setFillColor(CINZA)
    c.drawString(m + 4 * mm, ty - 3.2 * mm, 'ASSINATURA DO TRABALHADOR')
    c.drawString(m + larg * 0.68, ty - 3.2 * mm, 'DATA')
    y -= alt + 6 * mm

    # Tabela de entregas.
    fracs = [0.11, 0.37, 0.12, 0.07, 0.14, 0.19]
    xs = [m]
    for f in fracs:
        xs.append(xs[-1] + larg * f)
    h_cab = 8 * mm
    h_lin = 8.4 * mm
    c.setFillColor(TINTA)
    c.rect(m, y - h_cab, larg, h_cab, stroke=0, fill=1)
    c.setFillColor(HexColor('#ffffff'))
    c.setFont('Helvetica-Bold', 7.6)
    for i, nome in enumerate(COLUNAS):
        c.drawString(xs[i] + 1.8 * mm, y - h_cab + 2.9 * mm, nome)
    y -= h_cab
    c.setStrokeColor(LINHA)
    c.setLineWidth(0.5)
    topo = y
    for _ in range(LINHAS_PDF):
        y -= h_lin
        c.line(m, y, m + larg, y)
    for x in xs:
        c.line(x, topo, x, y)

    c.setFont('Helvetica', 7.2)
    c.setFillColor(CINZA)
    y -= 4.6 * mm
    c.drawString(m, y, 'Motivo: ' + ', '.join(MOTIVOS).lower().replace('primeira', 'Primeira', 1)
                 + '. Uma linha por item: o kit inteiro numa linha só não diz qual CA foi entregue.')

    # Rodapé.
    c.setFont('Helvetica', 6.8)
    yr = m - 2 * mm
    for ln in reversed(quebrar(c, rodape(dominio), 'Helvetica', 6.8, larg)):
        c.drawString(m, yr, ln)
        yr += 3 * mm
    c.showPage()
    c.save()


# ----------------------------------------------------------------- XLSX

def gerar_xlsx(termo, dominio, caminho):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Ficha de EPI'
    fino = Side(style='thin', color='9A9A9A')
    borda = Border(left=fino, right=fino, top=fino, bottom=fino)
    rotulo = Font(size=8, color='6B6B6B', bold=True)
    larguras = [12, 42, 14, 7, 18, 26]
    for i, w in enumerate(larguras):
        ws.column_dimensions[chr(65 + i)].width = w

    ws['A1'] = 'Ficha de controle de entrega de EPI'
    ws['A1'].font = Font(size=15, bold=True)
    ws.merge_cells('A1:F1')
    ws['A2'] = 'Registro do fornecimento de equipamento de proteção individual ao trabalhador.'
    ws['A2'].font = Font(size=9, color='6B6B6B')
    ws.merge_cells('A2:F2')

    # Cabeçalho: rótulo numa linha, espaço para preencher na de baixo.
    campos = [
        (4, [('A', 'F', 'EMPRESA')]),
        (6, [('A', 'D', 'TRABALHADOR'), ('E', 'F', 'MATRÍCULA')]),
        (8, [('A', 'B', 'FUNÇÃO'), ('C', 'F', 'SETOR')]),
        (10, [('A', 'A', 'ADMISSÃO'), ('B', 'B', 'Nº DO CALÇADO'), ('C', 'D', 'TAM. DA LUVA'), ('E', 'F', 'TAM. DA ROUPA')]),
    ]
    for linha, grupo in campos:
        for ini, fim, nome in grupo:
            ws[f'{ini}{linha}'] = nome
            ws[f'{ini}{linha}'].font = rotulo
            if ini != fim:
                ws.merge_cells(f'{ini}{linha}:{fim}{linha}')
                ws.merge_cells(f'{ini}{linha + 1}:{fim}{linha + 1}')
            ws[f'{ini}{linha + 1}'].border = Border(bottom=fino)
            for col in range(ord(ini), ord(fim) + 1):
                ws[f'{chr(col)}{linha + 1}'].border = Border(bottom=fino)
        ws.row_dimensions[linha + 1].height = 20

    ws['A13'] = 'Termo de recebimento e responsabilidade'
    ws['A13'].font = Font(size=10, bold=True)
    ws.merge_cells('A13:F13')
    ws['A14'] = termo
    ws['A14'].alignment = Alignment(wrap_text=True, vertical='top')
    ws['A14'].font = Font(size=9)
    ws['A14'].fill = PatternFill('solid', fgColor='F1EFEA')
    ws.merge_cells('A14:F14')
    ws.row_dimensions[14].height = 66
    ws['A16'] = 'ASSINATURA DO TRABALHADOR'
    ws['A16'].font = rotulo
    ws.merge_cells('A16:D16')
    ws['E16'] = 'DATA'
    ws['E16'].font = rotulo
    for col in 'ABCDEF':
        ws[f'{col}15'].border = Border(bottom=fino)
    ws.row_dimensions[15].height = 24

    cab = 18
    for i, nome in enumerate(COLUNAS):
        cel = ws.cell(row=cab, column=i + 1, value=nome)
        cel.font = Font(bold=True, color='FFFFFF', size=9)
        cel.fill = PatternFill('solid', fgColor='1A1A1A')
        cel.alignment = Alignment(vertical='center', wrap_text=True)
        cel.border = borda
    ws.row_dimensions[cab].height = 22
    for r in range(cab + 1, cab + 1 + LINHAS_XLSX):
        ws.row_dimensions[r].height = 20
        for col in range(1, 7):
            cel = ws.cell(row=r, column=col)
            cel.border = borda
            cel.alignment = Alignment(vertical='center', wrap_text=True)
        ws.cell(row=r, column=1).number_format = 'DD/MM/YYYY'

    dv = DataValidation(type='list', formula1='"' + ','.join(MOTIVOS) + '"', allow_blank=True)
    dv.error = 'Escolha um motivo da lista ou apague a célula.'
    dv.prompt = 'Uma linha por item, com o CA do item entregue.'
    ws.add_data_validation(dv)
    dv.add(f'E{cab + 1}:E{cab + LINHAS_XLSX}')

    fim = cab + LINHAS_XLSX + 2
    ws[f'A{fim}'] = rodape(dominio)
    ws[f'A{fim}'].font = Font(size=8, color='6B6B6B')
    ws[f'A{fim}'].alignment = Alignment(wrap_text=True, vertical='top')
    ws.merge_cells(f'A{fim}:F{fim}')
    ws.row_dimensions[fim].height = 34

    ws.print_title_rows = f'{cab}:{cab}'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = 'portrait'
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.4
    wb.properties.title = 'Ficha de controle de entrega de EPI'
    wb.properties.creator = dominio
    wb.save(caminho)


if __name__ == '__main__':
    termo, dominio = ler_fontes()
    os.makedirs(SAIDA, exist_ok=True)
    pdf = f'{SAIDA}/{NOME_ARQUIVO}.pdf'
    xlsx = f'{SAIDA}/{NOME_ARQUIVO}.xlsx'
    gerar_pdf(termo, dominio, pdf)
    gerar_xlsx(termo, dominio, xlsx)
    for f in (pdf, xlsx):
        print(f'{f}  {os.path.getsize(f) / 1024:.1f} KB')
